from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document

from dynamic_distillation.run_report_v1 import (
    _condition_rows,
    _first_and_last,
    _final_profile_assessment,
    _plot_balance_and_constraint_trends,
    _plot_steady_state_score,
    generate_core_v3_run_report,
    generate_run_report,
)


def _summary_rows() -> list[dict[str, float]]:
    common = {
        "F_lbmolph": 100.0,
        "P_bot_psia": 225.0,
        "T_Distillate_F": 110.0,
        "T_sump_F": 220.0,
        "Q_reb_used_BTUph": 5.0e6,
        "total_reflux_used_lbmolph": 75.0,
        "V_condensed_in_lbmolph": 90.0,
        "boilup_realized_lbmolph": 95.0,
        "Top_level_ctrl_sp": 0.5,
        "Bottom_level_ctrl_sp": 0.5,
        "Distillate_x_A": 0.9,
        "Distillate_x_B": 0.1,
        "Bottoms_x_A": 0.2,
        "Bottoms_x_B": 0.8,
    }
    return [
        common | {
            "time_s": 0.0,
            "D_lbmolph": 40.0,
            "B_lbmolph": 60.0,
            "P_top_psia": 224.0,
            "P_top_ctrl_pv_psia": 224.0,
            "P_top_psia_spec": 220.0,
            "Q_cond_used_BTUph": -4.0e6,
            "Top_level_ctrl_pv": 0.48,
            "Bottom_level_ctrl_pv": 0.45,
            "steady_state_score": 2.0,
            "steady_state_flag": 0.0,
        },
        common | {
            "time_s": 60.0,
            "D_lbmolph": 42.0,
            "B_lbmolph": 58.0,
            "P_top_psia": 220.2,
            "P_top_ctrl_pv_psia": 220.2,
            "P_top_psia_spec": 220.0,
            "Q_cond_used_BTUph": -4.4e6,
            "Top_level_ctrl_pv": 0.50,
            "Bottom_level_ctrl_pv": 0.49,
            "steady_state_score": 0.5,
            "steady_state_flag": 1.0,
        },
    ]


def _profile_rows() -> list[dict[str, float | str]]:
    rows = []
    for t in (0.0, 60.0):
        for stage in (1, 2, 3):
            rows.append(
                {
                    "time_s": t,
                    "stage": stage,
                    "node_type": "stage",
                    "T_F": 100.0 + 10.0 * stage,
                    "P_psia_hyd": 220.0 + stage,
                    "L_out_used_lbmolph": 70.0 + stage,
                    "V_out_lbmolph": 80.0 + stage,
                    "x_A": 0.9 - 0.2 * (stage - 1),
                    "x_B": 0.1 + 0.2 * (stage - 1),
                    "y_A": 0.95 - 0.2 * (stage - 1),
                    "y_B": 0.05 + 0.2 * (stage - 1),
                }
            )
    return rows


def test_first_and_last_orders_by_simulation_time() -> None:
    frame = pd.DataFrame([{"time_s": 10.0, "value": 2}, {"time_s": 0.0, "value": 1}])
    start, end = _first_and_last(frame)
    assert start["value"] == 1
    assert end["value"] == 2


def test_steady_state_score_chart_uses_history_and_acceptance_limit(tmp_path: Path) -> None:
    output = tmp_path / "steady_state_score.png"
    plotted = _plot_steady_state_score(
        pd.DataFrame(_summary_rows()), output, cumulative_offset_s=0.0,
    )
    assert plotted
    assert output.exists()
    assert output.stat().st_size > 1_000


def test_balance_and_constraint_chart_uses_persisted_evidence(tmp_path: Path) -> None:
    output = tmp_path / "balance_constraints.png"
    trajectory = pd.DataFrame([
        {"time_sec": 0.0, "solver_residual_inf_norm": 2.0e-9},
        {"time_sec": 10.0, "solver_residual_inf_norm": 3.0e-9},
    ])
    ledger = [
        {"time_sec": 0.0, "quantity": "A", "residual": 1.0},
        {"time_sec": 10.0, "quantity": "A", "residual": 2.0},
        {"time_sec": 0.0, "quantity": "energy", "residual": 3.0},
        {"time_sec": 10.0, "quantity": "energy", "residual": 4.0},
    ]
    assert _plot_balance_and_constraint_trends(
        trajectory, ledger, output, cumulative_offset_s=0.0,
    )
    assert output.exists()
    assert output.stat().st_size > 1_000


def test_final_profile_assessment_reports_tray_and_terminal_facts() -> None:
    rows = _profile_rows()
    rows.extend([
        {"time_s": 60.0, "stage": 0, "volume": "reflux_drum", "node_type": "reflux_drum", "T_F": 100.0, "P_psia_hyd": 220.0, "L_out_used_lbmolph": None, "V_out_lbmolph": None, "x_A": 0.9, "y_A": 0.9},
        {"time_s": 60.0, "stage": 4, "volume": "feed_tray", "node_type": "tray", "T_F": 130.0, "P_psia_hyd": 224.0, "L_out_used_lbmolph": -1.0, "V_out_lbmolph": 80.0, "x_A": 0.3, "y_A": 0.4},
    ])
    assessment = _final_profile_assessment(pd.DataFrame(rows))
    assert assessment["feed_stage"] == 4
    assert assessment["non_stage_nodes"] == ["reflux_drum"]
    assert assessment["top_to_bottom_pressure_drop_psia"] == 3.0
    assert "L_out_used_lbmolph:feed_tray" in assessment["flow_reversals"]
    assert assessment["composition_extrema"]["x_A"] == (0.3, 0.9)


def test_operating_snapshot_includes_declared_feed_conditions() -> None:
    rows = _condition_rows(
        {"F_lbmolph": 100.0, "Feed_temperature_F": 175.0, "Feed_pressure_psia": 232.06}
    )
    assert ["Feed", "100.00", "lbmol/h"] in rows
    assert ["Feed temperature", "175.00", "deg F"] in rows
    assert ["Feed pressure", "232.060", "psia"] in rows


def test_generate_run_report_creates_readable_docx(tmp_path: Path) -> None:
    summary = tmp_path / "column_summary_test.csv"
    profile = tmp_path / "column_profile_test.csv"
    metadata = tmp_path / "run_metadata_test.json"
    pd.DataFrame(_summary_rows()).to_csv(summary, index=False)
    pd.DataFrame(_profile_rows()).to_csv(profile, index=False)
    metadata.write_text(
        json.dumps(
            {
                "run_id": "test",
                "run_name": "C3/C4 report smoke test",
                "run_description": "Automated report validation",
                "status": "completed",
                "started_at_local": "2026-07-10 10:00:00",
                "ended_at_local": "2026-07-10 10:01:00",
                "elapsed_wall_sec": 30.0,
                "final_time_s": 60.0,
                "excel_path": str(tmp_path / "case.xlsx"),
                "summary_csv": str(summary),
                "profile_csv": str(profile),
                "native_checkpoint_init": {
                    "loaded": True,
                    "path": str(tmp_path / "seed.npz"),
                    "source_final_time_s": 900.0,
                },
            }
        ),
        encoding="utf-8",
    )

    report = Path(
        generate_run_report(
            metadata,
            simulation_parameters={"runtime_mode": "hydraulic", "thermo_mode": "dwsim", "n_steps": 300},
            launch_command="python -m dynamic_distillation.dynamic_run_scaffold_v1 --n-steps 300",
        )
    )

    assert report.exists()
    assert report.stat().st_size > 20_000
    doc = Document(report)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "C3/C4 report smoke test" in text
    assert "Operating Snapshot" in text
    assert "Final Tray Profiles" in text
    assert "Profile Summary" in text
    assert "Composition Extrema" in text
    assert "CLI Command" in text
    assert len(doc.inline_shapes) >= 6


def test_generate_run_report_accepts_current_core_v3_names_and_tray_nodes(tmp_path: Path) -> None:
    summary = tmp_path / "column_summary_current.csv"
    profile = tmp_path / "column_profile_current.csv"
    metadata = tmp_path / "run_metadata_current.json"
    rows = _summary_rows()
    for row in rows:
        row["Reflux_cmd_lbmolph"] = row.pop("total_reflux_used_lbmolph")
        row["Boilup_lbmolph"] = row.pop("boilup_realized_lbmolph")
        row["P_top_drum_psia"] = row.pop("P_top_ctrl_pv_psia")
        row["Bottom_level_fraction"] = row.pop("Bottom_level_ctrl_pv")
        row.pop("P_top_psia_spec")
    profile_rows = _profile_rows()
    for row in profile_rows:
        row["node_type"] = "tray"
    pd.DataFrame(rows).to_csv(summary, index=False)
    pd.DataFrame(profile_rows).to_csv(profile, index=False)
    metadata.write_text(
        json.dumps({
            "run_id": "current", "run_name": "Current Core V3 schema", "status": "completed",
            "summary_csv": str(summary), "profile_csv": str(profile),
        }),
        encoding="utf-8",
    )

    report = Path(generate_run_report(metadata))
    document = Document(report)
    final_table = next(
        table for table in document.tables
        if table.cell(0, 0).text == "Stage" and len(table.rows[0].cells) >= 8
    )
    assert len(final_table.rows) == 4
    assert final_table.cell(1, 0).text == "1"
    assert len(document.inline_shapes) >= 3


def test_core_v3_report_has_required_phase_one_sections(tmp_path: Path) -> None:
    output = tmp_path / "core_v3.docx"
    summary = {
        "schema_id": "core-v3-end-of-run-summary-v1", "time_sec": 60.0,
        "duties": {"condenser_BTUph": -100.0, "reboiler_BTUph": 120.0},
        "terminal_levels": {"distillate_drum_fraction": 0.5, "bottom_drum_fraction": 0.5},
        "products": {key: {"flow_lbmolph": 10.0, "temperature_F": 100.0, "pressure_psia": 15.0, "mole_fraction": {"A": 1.0}} for key in ("distillate", "bottoms")},
        "steady_state": {"steady": False, "score": 1.2, "terms": {"temperature_rate": 1.2}, "raw": {}},
        "profiles": [{"volume": "top", "node_type": "reflux_drum", "temperature_F": 100.0, "pressure_psia": 15.0, "liquid_inventory_lbmol": 2.0, "vapor_inventory_lbmol": 0.1, "liquid_flow_out_lbmolph": 10.0, "vapor_flow_out_lbmolph": None, "liquid_mole_fraction": {"A": 1.0}, "vapor_mole_fraction": {"A": 1.0}}],
    }
    generate_core_v3_run_report(summary, output_path=output, metadata={"pass_gate": True})
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for heading in ("Executive Verdict", "Run Narrative", "Initial, Final, and Delta", "Steady-State Assessment", "Operating-Limit Assessment", "Numerical Health", "Artifact Index"):
        assert heading in text
    assert len(document.inline_shapes) >= 1


def test_core_v3_docx_surfaces_controller_tuning_and_solver_history(tmp_path: Path) -> None:
    output = tmp_path / "core_v3_evidence.docx"
    trajectory = tmp_path / "accepted_steps.jsonl"
    ledger = tmp_path / "balance_ledger.jsonl"
    events = tmp_path / "events.jsonl"
    trajectory.write_text(
        "\n".join(json.dumps(row) for row in (
            {"time_sec": 0.0, "nfev": 8, "njev": 2, "retry_attempted": False, "controllers": {"drum_level": {"pv": 0.4, "sp": 0.5, "output_lbmolph": 10.0}}},
            {"time_sec": 1.0, "nfev": 10, "njev": 3, "retry_attempted": True, "controllers": {"drum_level": {"pv": 0.5, "sp": 0.5, "output_lbmolph": 11.0}}},
        )) + "\n",
        encoding="utf-8",
    )
    ledger.write_text(
        "\n".join(json.dumps(row) for row in (
            {"time_sec": 0.0, "volume": "global", "quantity": "A", "in": 10.0, "out": 9.0, "accumulation": 1.0, "residual": 0.0, "normalized_residual": 0.0},
            {"time_sec": 1.0, "volume": "global", "quantity": "A", "in": 10.0, "out": 9.0, "accumulation": 1.0, "residual": 1.0e-4, "normalized_residual": 1.0e-5},
            {"time_sec": 0.0, "volume": "global", "quantity": "energy", "net_expected_change": -20.0, "accumulation": -20.0, "residual": 0.0, "normalized_residual": 0.0},
            {"time_sec": 1.0, "volume": "global", "quantity": "energy", "net_expected_change": -20.0, "accumulation": -20.0, "residual": 2.0e-4, "normalized_residual": 1.0e-5},
        )) + "\n",
        encoding="utf-8",
    )
    events.write_text(
        json.dumps({"event_id": "restart", "time_sec": 0.0, "event_type": "checkpoint_restore", "variable": None, "before": None, "after": "seed.npz", "message": "Restored checkpoint.", "severity": "info", "source_component": "runner", "related_gate": None}) + "\n",
        encoding="utf-8",
    )
    summary = {
        "schema_id": "core-v3-end-of-run-summary-v1", "time_sec": 1.0,
        "duties": {"condenser_BTUph": -100.0, "reboiler_BTUph": 120.0},
        "terminal_levels": {"distillate_drum_fraction": 0.5, "bottom_drum_fraction": 0.5},
        "products": {key: {"flow_lbmolph": 10.0, "temperature_F": 100.0, "pressure_psia": 15.0, "mole_fraction": {"A": 1.0}} for key in ("distillate", "bottoms")},
        "steady_state": {"steady": False, "score": 1.2, "terms": {"temperature_rate": 1.2}, "raw": {}},
        "profiles": [{"volume": "top", "node_type": "reflux_drum", "temperature_F": 100.0, "pressure_psia": 15.0, "liquid_inventory_lbmol": 2.0, "vapor_inventory_lbmol": 0.1, "liquid_flow_out_lbmolph": 10.0, "vapor_flow_out_lbmolph": None, "liquid_mole_fraction": {"A": 1.0}, "vapor_mole_fraction": {"A": 1.0}}],
    }
    generate_core_v3_run_report(
        summary,
        output_path=output,
        metadata={
            "report_trajectory_jsonl": trajectory,
            "report_balance_ledger_jsonl": ledger,
            "report_events_jsonl": events,
            "controller_tuning": {"drum_kc": 4.0, "drum_ti_sec": 120.0},
        },
    )
    document = Document(output)
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    assert "drum_kc" in text
    assert "Solver and Acceptance" in text
    assert "Function evaluations total / max" in text
    assert "Runtime and Cache" in text
    assert "Initial accepted interval" in text
    assert "Final accepted interval" in text
    assert "Material Balance" in text
    assert "Energy Balance" in text
    assert "Balance Summary" in text
    assert "Integrated |residual|" in text
    assert "Net expected" in text
    assert "Tracking Performance" in text
    assert "Output and Settling" in text
    assert "Mean |error|" in text
    assert "Overshoot" in text
    assert "At low bound (s)" in text
    assert "Time (s)" in text
    assert "Related gate" in text


def test_core_v3_docx_splits_three_component_product_composition(tmp_path: Path) -> None:
    output = tmp_path / "three_component.docx"
    composition = {"A": 0.2, "B": 0.3, "C": 0.5}
    summary = {
        "schema_id": "core-v3-end-of-run-summary-v1", "time_sec": 1.0,
        "duties": {"condenser_BTUph": -100.0, "reboiler_BTUph": 120.0},
        "terminal_levels": {"distillate_drum_fraction": 0.5, "bottom_drum_fraction": 0.5},
        "products": {key: {"flow_lbmolph": 10.0, "temperature_F": 100.0, "pressure_psia": 15.0, "mole_fraction": composition} for key in ("distillate", "bottoms")},
        "steady_state": {"steady": True, "score": 0.0, "terms": {}, "raw": {}},
        "profiles": [{"volume": "top", "node_type": "reflux_drum", "temperature_F": 100.0, "pressure_psia": 15.0, "liquid_inventory_lbmol": 2.0, "vapor_inventory_lbmol": 0.1, "liquid_flow_out_lbmolph": 10.0, "vapor_flow_out_lbmolph": None, "liquid_mole_fraction": composition, "vapor_mole_fraction": composition}],
    }
    generate_core_v3_run_report(summary, output_path=output, metadata={"pass_gate": True})
    document = Document(output)
    assert "Product Composition" in "\n".join(p.text for p in document.paragraphs)
    assert any(
        [cell.text for cell in row.cells] == ["Distillate", "A", "0.200000"]
        for table in document.tables for row in table.rows
    )
